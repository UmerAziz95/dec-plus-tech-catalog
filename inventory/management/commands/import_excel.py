"""
Management command to import car parts data from the client's Excel file.

Expected file: Mitsubishi_database_Id_step1-2-3 sample.xlsx
  - Sheet 'CarId': CarId, Car Model, Steering, Transmission, WD, Engine, Car parameters, Additional Note
  - Sheet 'carId & GroupId': 3 repeated column pairs (carId, GroupID) separated by empty columns
  - Sheet 'groupId & part_number': 3 repeated column sets (groupId, Brand, part_number) separated by empty columns

Usage:
    python manage.py import_excel "path/to/file.xlsx"
    python manage.py import_excel "path/to/file.xlsx" --clear   # Clear existing data first
    python manage.py import_excel "path/to/file.xlsx" --sheet cars  # Import only cars
"""
import time
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from inventory.models import Car, Group, CarGroup, Part


class Command(BaseCommand):
    help = 'Import car parts data from the client Excel file into the database.'

    def add_arguments(self, parser):
        parser.add_argument(
            'filepath',
            type=str,
            help='Path to the Excel file (.xlsx)'
        )
        parser.add_argument(
            '--clear',
            action='store_true',
            help='Clear all existing inventory data before importing'
        )
        parser.add_argument(
            '--sheet',
            type=str,
            choices=['cars', 'groups', 'parts', 'all'],
            default='all',
            help='Which sheet(s) to import (default: all)'
        )
        parser.add_argument(
            '--batch-size',
            type=int,
            default=5000,
            help='Batch size for bulk operations (default: 5000)'
        )

    def handle(self, *args, **options):
        filepath = options['filepath']
        clear = options['clear']
        sheet = options['sheet']
        batch_size = options['batch_size']

        try:
            import openpyxl
        except ImportError:
            raise CommandError('openpyxl is required. Install it: pip install openpyxl')

        self.stdout.write(self.style.NOTICE(f'Opening file: {filepath}'))

        try:
            wb = openpyxl.load_workbook(filepath, read_only=True)
        except FileNotFoundError:
            raise CommandError(f'File not found: {filepath}')
        except Exception as e:
            raise CommandError(f'Error opening file: {e}')

        self.stdout.write(f'Sheets found: {wb.sheetnames}')

        if clear:
            self.stdout.write(self.style.WARNING('Clearing existing data...'))
            Part.objects.all().delete()
            CarGroup.objects.all().delete()
            Group.objects.all().delete()
            Car.objects.all().delete()
            self.stdout.write(self.style.SUCCESS('Existing data cleared.'))

        total_start = time.time()

        try:
            if sheet in ('cars', 'all'):
                self._import_cars(wb, batch_size)

            if sheet in ('groups', 'all'):
                self._import_car_groups(wb, batch_size)

            if sheet in ('parts', 'all'):
                self._import_parts(wb, batch_size)
        finally:
            wb.close()

        total_time = time.time() - total_start
        self.stdout.write(self.style.SUCCESS(
            f'\nImport completed in {total_time:.1f} seconds!'
        ))
        self._print_summary()

    def _import_cars(self, wb, batch_size):
        """Import cars from the 'CarId' sheet."""
        sheet_name = 'CarId'
        if sheet_name not in wb.sheetnames:
            raise CommandError(f"Sheet '{sheet_name}' not found in workbook.")

        self.stdout.write(self.style.NOTICE(f'\n--- Importing Cars from "{sheet_name}" ---'))
        ws = wb[sheet_name]
        start = time.time()

        cars_to_create = []
        skipped = 0
        processed = 0

        # Get existing car_ids to skip duplicates
        existing_ids = set(Car.objects.values_list('car_id', flat=True))

        for i, row in enumerate(ws.iter_rows(min_row=2)):
            values = [cell.value for cell in row]
            car_id = values[0]

            if not car_id:
                continue

            car_id = str(car_id).strip()

            if car_id in existing_ids:
                skipped += 1
                continue

            cars_to_create.append(Car(
                car_id=car_id,
                car_model=str(values[1] or '').strip(),
                steering=str(values[2] or '').strip(),
                transmission=str(values[3] or '').strip(),
                wd=str(values[4] or '').strip(),
                engine=str(values[5] or '').strip(),
                car_parameters=str(values[6] or '').strip(),
                additional_note=str(values[7] or '').strip() if len(values) > 7 else '',
            ))
            existing_ids.add(car_id)
            processed += 1

            if len(cars_to_create) >= batch_size:
                Car.objects.bulk_create(cars_to_create, ignore_conflicts=True)
                self.stdout.write(f'  Cars: {processed} processed...')
                cars_to_create = []

        if cars_to_create:
            Car.objects.bulk_create(cars_to_create, ignore_conflicts=True)

        elapsed = time.time() - start
        self.stdout.write(self.style.SUCCESS(
            f'  Cars: {processed} imported, {skipped} skipped (duplicates) in {elapsed:.1f}s'
        ))

    def _import_car_groups(self, wb, batch_size):
        """Import car-group relationships from 'carId & GroupId' sheet."""
        sheet_name = 'carId & GroupId'
        if sheet_name not in wb.sheetnames:
            raise CommandError(f"Sheet '{sheet_name}' not found in workbook.")

        self.stdout.write(self.style.NOTICE(f'\n--- Importing Car-Group links from "{sheet_name}" ---'))
        ws = wb[sheet_name]
        start = time.time()

        # Build lookup of existing cars
        car_lookup = dict(Car.objects.values_list('car_id', 'id'))
        self.stdout.write(f'  Loaded {len(car_lookup)} cars for lookup')

        # Phase 1: Collect all unique group_ids and car-group pairs
        self.stdout.write('  Phase 1: Reading sheet to collect groups and pairs...')
        all_group_ids = set()
        all_pairs = []  # (car_id_str, group_id_str)
        row_count = 0

        for row in ws.iter_rows(min_row=2):
            values = [cell.value for cell in row]
            # 3 column pairs: (0,1), (3,4), (6,7)
            for col_pair in [(0, 1), (3, 4), (6, 7)]:
                car_id = values[col_pair[0]] if len(values) > col_pair[0] else None
                group_id = values[col_pair[1]] if len(values) > col_pair[1] else None

                if car_id and group_id:
                    car_id_str = str(car_id).strip()
                    group_id_str = str(group_id).strip()
                    all_group_ids.add(group_id_str)
                    all_pairs.append((car_id_str, group_id_str))

            row_count += 1
            if row_count % 200000 == 0:
                self.stdout.write(f'    Read {row_count} rows, {len(all_pairs)} pairs so far...')

        self.stdout.write(f'  Found {len(all_group_ids)} unique groups, {len(all_pairs)} total pairs')

        # Phase 2: Create Group objects
        self.stdout.write('  Phase 2: Creating Group objects...')
        existing_group_ids = set(Group.objects.values_list('group_id', flat=True))
        new_group_ids = all_group_ids - existing_group_ids

        groups_to_create = [Group(group_id=gid) for gid in new_group_ids]
        created_count = 0
        for i in range(0, len(groups_to_create), batch_size):
            batch = groups_to_create[i:i + batch_size]
            Group.objects.bulk_create(batch, ignore_conflicts=True)
            created_count += len(batch)
            if created_count % 50000 == 0:
                self.stdout.write(f'    Created {created_count} groups...')

        self.stdout.write(self.style.SUCCESS(f'  Groups: {len(new_group_ids)} new created'))

        # Phase 3: Create CarGroup links
        self.stdout.write('  Phase 3: Creating Car-Group links...')
        group_lookup = dict(Group.objects.values_list('group_id', 'id'))

        # Deduplicate pairs
        unique_pairs = set(all_pairs)
        self.stdout.write(f'  {len(unique_pairs)} unique pairs to link')

        # Get existing links to skip
        existing_links = set(
            CarGroup.objects.values_list('car_id', 'group_id')
        )

        links_to_create = []
        skipped = 0
        errors = 0
        created = 0

        for car_id_str, group_id_str in unique_pairs:
            car_pk = car_lookup.get(car_id_str)
            group_pk = group_lookup.get(group_id_str)

            if not car_pk or not group_pk:
                errors += 1
                continue

            if (car_pk, group_pk) in existing_links:
                skipped += 1
                continue

            links_to_create.append(CarGroup(car_id=car_pk, group_id=group_pk))
            existing_links.add((car_pk, group_pk))
            created += 1

            if len(links_to_create) >= batch_size:
                CarGroup.objects.bulk_create(links_to_create, ignore_conflicts=True)
                self.stdout.write(f'    Created {created} links...')
                links_to_create = []

        if links_to_create:
            CarGroup.objects.bulk_create(links_to_create, ignore_conflicts=True)

        elapsed = time.time() - start
        self.stdout.write(self.style.SUCCESS(
            f'  Car-Group links: {created} created, {skipped} skipped, {errors} errors in {elapsed:.1f}s'
        ))

    def _import_parts(self, wb, batch_size):
        """Import parts from 'groupId & part_number' sheet."""
        sheet_name = 'groupId & part_number'
        if sheet_name not in wb.sheetnames:
            raise CommandError(f"Sheet '{sheet_name}' not found in workbook.")

        self.stdout.write(self.style.NOTICE(f'\n--- Importing Parts from "{sheet_name}" ---'))
        ws = wb[sheet_name]
        start = time.time()

        # Build lookup
        group_lookup = dict(Group.objects.values_list('group_id', 'id'))
        self.stdout.write(f'  Loaded {len(group_lookup)} groups for lookup')

        # Collect existing parts to avoid duplicates
        self.stdout.write('  Loading existing parts for dedup...')
        existing_parts = set(
            Part.objects.values_list('group_id', 'part_number')
        )
        self.stdout.write(f'  {len(existing_parts)} existing parts loaded')

        parts_to_create = []
        skipped = 0
        errors = 0
        created = 0
        row_count = 0

        for row in ws.iter_rows(min_row=2):
            values = [cell.value for cell in row]
            # 3 column sets: (0,1,2), (4,5,6), (8,9,10)
            for col_set in [(0, 1, 2), (4, 5, 6), (8, 9, 10)]:
                gid = values[col_set[0]] if len(values) > col_set[0] else None
                brand = values[col_set[1]] if len(values) > col_set[1] else None
                pn = values[col_set[2]] if len(values) > col_set[2] else None

                if not gid or not pn:
                    continue

                gid_str = str(gid).strip()
                pn_str = str(pn).strip()
                brand_str = str(brand or '').strip()

                group_pk = group_lookup.get(gid_str)
                if not group_pk:
                    errors += 1
                    continue

                if (group_pk, pn_str) in existing_parts:
                    skipped += 1
                    continue

                parts_to_create.append(Part(
                    group_id=group_pk,
                    brand=brand_str,
                    part_number=pn_str,
                ))
                existing_parts.add((group_pk, pn_str))
                created += 1

                if len(parts_to_create) >= batch_size:
                    Part.objects.bulk_create(parts_to_create, ignore_conflicts=True)
                    self.stdout.write(f'    Parts: {created} created...')
                    parts_to_create = []

            row_count += 1
            if row_count % 200000 == 0:
                self.stdout.write(f'    Processed {row_count} rows...')

        if parts_to_create:
            Part.objects.bulk_create(parts_to_create, ignore_conflicts=True)

        elapsed = time.time() - start
        self.stdout.write(self.style.SUCCESS(
            f'  Parts: {created} created, {skipped} skipped, {errors} errors in {elapsed:.1f}s'
        ))

    def _print_summary(self):
        """Print a summary of the current database state."""
        self.stdout.write(self.style.NOTICE('\n--- Database Summary ---'))
        self.stdout.write(f'  Cars:           {Car.objects.count():,}')
        self.stdout.write(f'  Groups:         {Group.objects.count():,}')
        self.stdout.write(f'  Car-Group links: {CarGroup.objects.count():,}')
        self.stdout.write(f'  Parts:          {Part.objects.count():,}')
