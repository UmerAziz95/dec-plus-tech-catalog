from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.test import Client, SimpleTestCase, TestCase
from django.urls import reverse

from inventory.models import BasketCrossCode, Car, CarGroup, ImportBatch, Part, PartCrossCode
from django.core.files.uploadedfile import SimpleUploadedFile
from inventory.services.basket_service import BasketServiceCrossCode
from inventory.services.bulk_search_service import BulkSearchService, BulkSearchServiceCrossCode
from inventory.services.import_services import ExcelImportService, ImportBatchDeleteService, normalize_crosscode_product_brand
from inventory.services.manual_entry_service import ManualEntryService
from inventory.services.part_number_utils import (
    crosscode_wildcard_to_like,
    sanitize_crosscode_bulk_part,
    sanitize_crosscode_search,
    sanitize_part_number,
)
from inventory.services.part_search_service import (
    SOURCE_CATALOG_IMPORT,
    SOURCE_CAR_WITH_PARTS,
    SOURCE_MANUAL_ENTRY,
    PartSearchService,
    PartSearchServiceCrossCode,
)

User = get_user_model()


class PartNumberSanitizationTests(SimpleTestCase):
    def test_removes_documented_special_characters(self):
        self.assertEqual(sanitize_part_number('929-00*0=803()21+B'), '92900080321B')

    def test_strips_whitespace(self):
        self.assertEqual(sanitize_part_number('  MQ900871  '), 'MQ900871')

    def test_returns_empty_for_only_special_characters(self):
        self.assertEqual(sanitize_part_number(')(.,:;/?\\|*+-=@!%&#$_'), '')

    def test_strips_at_hash_dollar_symbols(self):
        self.assertEqual(sanitize_part_number('MQ@#$901069'), 'MQ901069')

    def test_strips_underscore(self):
        self.assertEqual(sanitize_part_number('MR_418807'), 'MR418807')

    def test_handles_none(self):
        self.assertEqual(sanitize_part_number(None), '')


class CrossCodeSanitizationTests(SimpleTestCase):
    def test_interactive_strips_symbols_and_spaces_keeps_star(self):
        self.assertEqual(
            sanitize_crosscode_search('43 022*!@#A01', keep_star=True),
            '43022*A01',
        )

    def test_interactive_uppercases(self):
        self.assertEqual(sanitize_crosscode_search('d1086'), 'D1086')

    def test_bulk_removes_star(self):
        cleaned, had_star = sanitize_crosscode_bulk_part('43022*A01')
        self.assertTrue(had_star)
        self.assertEqual(cleaned, '43022A01')

    def test_bulk_no_star(self):
        cleaned, had_star = sanitize_crosscode_bulk_part('d-1086')
        self.assertFalse(had_star)
        self.assertEqual(cleaned, 'D1086')

    def test_wildcard_to_like(self):
        self.assertEqual(crosscode_wildcard_to_like('43022*A01'), '43022%A01')

    def test_jikiu_catalogue_url_becomes_product_brand_name(self):
        self.assertEqual(
            normalize_crosscode_product_brand('https://www.jikiu.jp/catalogue/804628'),
            'JIKIU',
        )
        self.assertEqual(normalize_crosscode_product_brand('NIBK'), 'NIBK')


class CrossCodeSearchRulesTests(TestCase):
    def setUp(self):
        PartCrossCode.objects.create(
            brand='NIBK', product_no='P1', oe_brand='MIT', part_number='D1086', group_id='P1',
        )
        PartCrossCode.objects.create(
            brand='NIBK', product_no='P2', oe_brand='MIT', part_number='D10867418', group_id='P2',
        )
        PartCrossCode.objects.create(
            brand='NIBK', product_no='P3', oe_brand='MIT', part_number='MR955727', group_id='P3',
        )
        PartCrossCode.objects.create(
            brand='NIBK', product_no='P4', oe_brand='MIT', part_number='43022XA01', group_id='P4',
        )
        PartCrossCode.objects.create(
            brand='NIBK', product_no='P5', oe_brand='MIT', part_number='43022YA01', group_id='P5',
        )

    def test_unique_exact_skips_did_you_mean_and_opens_family(self):
        needs, candidates, query = PartSearchServiceCrossCode.needs_disambiguation('MR955727')
        self.assertFalse(needs)
        self.assertEqual(candidates, [])
        self.assertEqual(sanitize_part_number(query), 'MR955727')

    def test_same_code_same_family_different_brands_asks_did_you_mean(self):
        PartCrossCode.objects.create(
            brand='NIBK', product_no='FAM', oe_brand='TOYOTA', part_number='SHARE1', group_id='FAM',
        )
        PartCrossCode.objects.create(
            brand='NIBK', product_no='FAM', oe_brand='HONDA', part_number='SHARE1', group_id='FAM',
        )
        needs, candidates, query = PartSearchServiceCrossCode.needs_disambiguation('SHARE1')
        self.assertTrue(needs)
        pairs = {
            (c['oe_brand'], sanitize_part_number(c['part_number']))
            for c in candidates
        }
        self.assertEqual(pairs, {('TOYOTA', 'SHARE1'), ('HONDA', 'SHARE1')})

    def test_same_code_different_brands_asks_did_you_mean(self):
        PartCrossCode.objects.create(
            brand='NIBK', product_no='DUP1', oe_brand='TOYOTA', part_number='SAMECODE1', group_id='DUP1',
        )
        PartCrossCode.objects.create(
            brand='NIBK', product_no='DUP2', oe_brand='HONDA', part_number='SAMECODE1', group_id='DUP2',
        )
        needs, candidates, query = PartSearchServiceCrossCode.needs_disambiguation('SAMECODE1')
        self.assertTrue(needs)
        self.assertEqual(sanitize_part_number(query), 'SAMECODE1')
        pairs = {
            (c['oe_brand'], sanitize_part_number(c['part_number']))
            for c in candidates
        }
        self.assertEqual(pairs, {('TOYOTA', 'SAMECODE1'), ('HONDA', 'SAMECODE1')})

    def test_bulk_candidates_are_exact_matches_only(self):
        rows = [
            {'brand': 'A', 'brand_number': 'B', 'part_number': 'D1086', 'bulk_miss': False},
            {'brand': 'A', 'brand_number': 'B', 'part_number': '43022*A01', 'bulk_miss': True},
            {'brand': 'A', 'brand_number': 'B', 'part_number': 'D10867418', 'bulk_miss': False},
        ]
        candidates = PartSearchServiceCrossCode.candidates_from_bulk_rows(rows)
        norms = {sanitize_part_number(c['part_number']) for c in candidates}
        self.assertEqual(norms, {'D1086', 'D10867418'})
        self.assertTrue(all(c.get('oe_brand') for c in candidates))

    def test_prefix_asks_did_you_mean_when_no_exact_code(self):
        needs, candidates, query = PartSearchServiceCrossCode.needs_disambiguation('43022')
        self.assertTrue(needs)
        self.assertEqual(query, '43022')
        norms = {sanitize_part_number(c['part_number']) for c in candidates}
        self.assertIn('43022XA01', norms)
        self.assertIn('43022YA01', norms)
        self.assertTrue(all(c.get('oe_brand') == 'MIT' for c in candidates))

    def test_wildcard_asks_did_you_mean(self):
        needs, candidates, query = PartSearchServiceCrossCode.needs_disambiguation('43022*A01')
        self.assertTrue(needs)
        self.assertEqual(query, '43022*A01')
        norms = {sanitize_part_number(c['part_number']) for c in candidates}
        self.assertIn('43022XA01', norms)
        self.assertIn('43022YA01', norms)
        self.assertTrue(all('oe_brand' in c and 'part_number' in c for c in candidates))

    def test_wildcard_brand_name_asks_did_you_mean(self):
        needs, candidates, query = PartSearchServiceCrossCode.needs_disambiguation('MI*')
        self.assertTrue(needs)
        self.assertEqual(query, 'MI*')
        self.assertTrue(all(c.get('oe_brand') == 'MIT' for c in candidates))
        norms = {sanitize_part_number(c['part_number']) for c in candidates}
        self.assertIn('D1086', norms)
        self.assertIn('43022XA01', norms)

    def test_product_no_prefix_does_not_list_unrelated_family_codes(self):
        PartCrossCode.objects.create(
            brand='NIBK', product_no='PN9001', oe_brand='TOYOTA', part_number='CODE9001A', group_id='PN9001',
        )
        PartCrossCode.objects.create(
            brand='NIBK', product_no='PN9001', oe_brand='HONDA', part_number='CODE9001B', group_id='PN9001',
        )
        candidates = PartSearchServiceCrossCode.find_candidate_codes('PN9001')
        norms = {sanitize_part_number(c['part_number']) for c in candidates}
        self.assertNotIn('CODE9001A', norms)
        self.assertNotIn('CODE9001B', norms)

    def test_prefix_did_you_mean_only_lists_matching_codes(self):
        PartCrossCode.objects.create(
            brand='JIKIU', product_no='3PK1000', oe_brand='BANDO',
            part_number='3PK1000', group_id='3PK1000',
        )
        PartCrossCode.objects.create(
            brand='JIKIU', product_no='3PK1000', oe_brand='HONDA',
            part_number='31110PH7013', group_id='3PK1000',
        )
        PartCrossCode.objects.create(
            brand='JIKIU', product_no='3PK1000', oe_brand='MAZDA',
            part_number='N39013813', group_id='3PK1000',
        )
        PartCrossCode.objects.create(
            brand='JIKIU', product_no='3PK1000', oe_brand='MITSUBISHI',
            part_number='MD361163', group_id='3PK1000',
        )
        PartCrossCode.objects.create(
            brand='JIKIU', product_no='3PK1055', oe_brand='DAYCO',
            part_number='3PK1055', group_id='3PK1055',
        )
        needs, candidates, query = PartSearchServiceCrossCode.needs_disambiguation('3PK10')
        self.assertTrue(needs)
        self.assertEqual(query, '3PK10')
        norms = {sanitize_part_number(c['part_number']) for c in candidates}
        self.assertEqual(norms, {'3PK1000', '3PK1055'})
        pairs = {(c['oe_brand'], sanitize_part_number(c['part_number'])) for c in candidates}
        self.assertEqual(pairs, {('BANDO', '3PK1000'), ('DAYCO', '3PK1055')})

    def test_wildcard_did_you_mean_only_lists_matching_codes(self):
        PartCrossCode.objects.create(
            brand='JIKIU', product_no='3PK1000', oe_brand='BANDO',
            part_number='3PK1000', group_id='3PK1000',
        )
        PartCrossCode.objects.create(
            brand='JIKIU', product_no='3PK1000', oe_brand='HONDA',
            part_number='31110PH7013', group_id='3PK1000',
        )
        PartCrossCode.objects.create(
            brand='JIKIU', product_no='3PK1000', oe_brand='MITSUBISHI',
            part_number='4451A022', group_id='3PK1000',
        )
        PartCrossCode.objects.create(
            brand='JIKIU', product_no='3PK1055', oe_brand='DAYCO',
            part_number='3PK1055', group_id='3PK1055',
        )
        needs, candidates, query = PartSearchServiceCrossCode.needs_disambiguation('3PK10*')
        self.assertTrue(needs)
        self.assertEqual(query, '3PK10*')
        norms = {sanitize_part_number(c['part_number']) for c in candidates}
        self.assertEqual(norms, {'3PK1000', '3PK1055'})
        self.assertNotIn('31110PH7013', norms)
        self.assertNotIn('4451A022', norms)

    def test_exact_product_no_skips_did_you_mean_and_opens_family(self):
        PartCrossCode.objects.create(
            brand='NIBK', product_no='PN9001', oe_brand='TOYOTA', part_number='CODE9001A', group_id='PN9001',
        )
        PartCrossCode.objects.create(
            brand='NIBK', product_no='PN9001', oe_brand='HONDA', part_number='CODE9001B', group_id='PN9001',
        )
        needs, candidates, query = PartSearchServiceCrossCode.needs_disambiguation('PN9001')
        self.assertFalse(needs)
        self.assertEqual(candidates, [])
        self.assertEqual(query, 'PN9001')

        results = PartSearchServiceCrossCode.build_results('PN9001')
        codes = {sanitize_part_number(item['part'].part_number) for item in results}
        self.assertEqual(codes, {'CODE9001A', 'CODE9001B'})

    def test_exact_code_with_longer_sibling_still_opens_family(self):
        needs, candidates, query = PartSearchServiceCrossCode.needs_disambiguation('D1086')
        self.assertFalse(needs)
        self.assertEqual(candidates, [])
        self.assertEqual(sanitize_part_number(query), 'D1086')

    def test_exact_search_after_confirm(self):
        results = PartSearchServiceCrossCode.build_results('D1086')
        self.assertEqual(len(results), 1)
        self.assertEqual(sanitize_part_number(results[0]['part'].part_number), 'D1086')

        results_long = PartSearchServiceCrossCode.build_results('D10867418')
        self.assertEqual(len(results_long), 1)
        self.assertEqual(sanitize_part_number(results_long[0]['part'].part_number), 'D10867418')

    def test_bulk_misses_star_rows_and_exact_matches_only(self):
        rows = [
            {'brand': 'A', 'brand_number': 'B', 'part_number': 'D1086', 'bulk_miss': False},
            {'brand': 'A', 'brand_number': 'B', 'part_number': '43022*A01', 'bulk_miss': True},
            {'brand': 'A', 'brand_number': 'B', 'part_number': 'D10867418', 'bulk_miss': False},
        ]
        # Simulate parse output for star row already flagged; also run through bulk helper path
        cleaned, had_star = sanitize_crosscode_bulk_part('43022*A01')
        self.assertTrue(had_star)

        results, summary, err = BulkSearchServiceCrossCode.run_bulk_search(
            user=None,
            rows_data=rows,
            save_to_basket=False,
        )
        self.assertIsNone(err)
        statuses = {r['part_number']: r['status'] for r in results}
        self.assertEqual(statuses['D1086'], 'Found')
        self.assertEqual(statuses['D10867418'], 'Found')
        self.assertEqual(statuses['43022*A01'], 'Missed (*)')
        self.assertEqual(summary['found_count'], 2)
        self.assertEqual(summary['missed_count'], 1)
        found_d1086 = next(r for r in results if r['part_number'] == 'D1086')
        self.assertTrue(found_d1086['match_pairs'])
        self.assertEqual(found_d1086['match_pairs'][0]['product_no'], 'P1')
        self.assertEqual(found_d1086['match_pairs'][0]['brand'], 'MIT')
        self.assertEqual(
            sanitize_part_number(found_d1086['match_pairs'][0]['code']),
            'D1086',
        )

    def test_family_results_can_filter_to_selected_brand(self):
        PartCrossCode.objects.create(
            brand='NIBK', product_no='FAM1', oe_brand='TOYOTA', part_number='CODE230020', group_id='FAM1',
        )
        PartCrossCode.objects.create(
            brand='ATE', product_no='FAM2', oe_brand='HONDA', part_number='CODE230020', group_id='FAM2',
        )
        toyota_family = PartSearchServiceCrossCode.build_results('CODE230020', oe_brand='TOYOTA')
        honda_family = PartSearchServiceCrossCode.build_results('CODE230020', oe_brand='HONDA')
        self.assertTrue(toyota_family)
        self.assertTrue(honda_family)
        self.assertTrue(all(item['part'].product_no == 'FAM1' for item in toyota_family))
        self.assertTrue(all(item['part'].product_no == 'FAM2' for item in honda_family))

    def test_same_code_same_family_both_brands_open_family(self):
        PartCrossCode.objects.create(
            brand='NIBK', product_no='PN3469', oe_brand='CHRYSLER',
            part_number='MR955727', group_id='PN3469',
        )
        PartCrossCode.objects.create(
            brand='NIBK', product_no='PN3469', oe_brand='MITSUBISHI',
            part_number='MR955727', group_id='PN3469',
        )
        PartCrossCode.objects.create(
            brand='NIBK', product_no='PN3469', oe_brand='BENDIX',
            part_number='DB1441', group_id='PN3469',
        )
        chrysler_family = PartSearchServiceCrossCode.build_results(
            'MR955727', oe_brand='CHRYSLER',
        )
        mitsu_family = PartSearchServiceCrossCode.build_results(
            'MR955727', oe_brand='MITSUBISHI',
        )
        self.assertTrue(chrysler_family)
        self.assertTrue(mitsu_family)
        chrysler_codes = {
            sanitize_part_number(item['part'].part_number) for item in chrysler_family
        }
        mitsu_codes = {
            sanitize_part_number(item['part'].part_number) for item in mitsu_family
        }
        self.assertEqual(chrysler_codes, {'MR955727', 'DB1441'})
        self.assertEqual(mitsu_codes, {'MR955727', 'DB1441'})
        self.assertTrue(all(item['part'].product_no == 'PN3469' for item in chrysler_family))
        self.assertTrue(all(item['part'].product_no == 'PN3469' for item in mitsu_family))

    def test_single_search_results_include_brand_and_code(self):
        results = PartSearchServiceCrossCode.build_results('D1086')
        self.assertEqual(len(results), 1)
        part = results[0]['part']
        self.assertEqual(part.oe_brand, 'MIT')
        self.assertEqual(sanitize_part_number(part.part_number), 'D1086')

    def test_family_pagination_keeps_exact_flag(self):
        user = User.objects.create_user(email='page@example.com', password='Test@123')
        client = Client()
        client.force_login(user)
        for i in range(26):
            PartCrossCode.objects.create(
                brand='NIBK',
                product_no='P3',
                oe_brand=f'BR{i:02d}',
                part_number=f'MR955727X{i:02d}',
                group_id='P3',
            )
        family_url = reverse('inventory:search_part_crosscode')
        page_two = client.get(family_url, {'q': 'MR955727', 'exact': '1', 'page': '2'})
        self.assertNotContains(page_two, 'id="did-you-mean-table"')
        self.assertContains(page_two, 'exact=1')
        self.assertContains(page_two, 'search-results-table')

        dropped_exact = client.get(family_url, {'q': 'MR955727', 'page': '2'})
        self.assertNotContains(dropped_exact, 'id="did-you-mean-table"')
        self.assertContains(dropped_exact, 'search-results-table')


class CrossCodePartsCatalogSearchTests(TestCase):
    def setUp(self):
        PartCrossCode.objects.create(
            brand='NIBK', product_no='P4', oe_brand='MIT', part_number='43022XA01', group_id='P4',
        )
        PartCrossCode.objects.create(
            brand='NIBK', product_no='P5', oe_brand='MIT', part_number='43022YA01', group_id='P5',
        )
        PartCrossCode.objects.create(
            brand='NIBK', product_no='P1', oe_brand='MIT', part_number='D1086', group_id='P1',
        )

    def test_wildcard_matches_code_like_search_by_part_number(self):
        from inventory.services.parts_catalog_service import PartsCatalogServiceCrossCode
        codes = set(
            PartsCatalogServiceCrossCode.build_queryset('43022*A01').values_list('part_number', flat=True)
        )
        self.assertEqual(codes, {'43022XA01', '43022YA01'})

    def test_plain_search_still_matches_brand_and_code(self):
        from inventory.services.parts_catalog_service import PartsCatalogServiceCrossCode
        codes = set(
            PartsCatalogServiceCrossCode.build_queryset('D1086').values_list('part_number', flat=True)
        )
        self.assertEqual(codes, {'D1086'})

    def test_wildcard_matches_brand_name(self):
        from inventory.services.parts_catalog_service import PartsCatalogServiceCrossCode
        codes = set(
            PartsCatalogServiceCrossCode.build_queryset('MI*').values_list('part_number', flat=True)
        )
        self.assertEqual(codes, {'43022XA01', '43022YA01', 'D1086'})

    def test_suggestions_prefix_match_codes(self):
        from inventory.services.parts_catalog_service import PartsCatalogServiceCrossCode
        self.assertEqual(
            PartsCatalogServiceCrossCode.suggest_values('43022'),
            ['43022XA01', '43022YA01'],
        )


class CrossCodeBrandNamesCombinedTests(TestCase):
    def setUp(self):
        PartCrossCode.objects.create(
            brand='NIBK', product_no='P1', oe_brand='HONDA', part_number='C1', group_id='P1',
        )
        PartCrossCode.objects.create(
            brand='NIBK', product_no='P2', oe_brand='TOYOTA', part_number='C2', group_id='P2',
        )
        PartCrossCode.objects.create(
            brand='ATE', product_no='P3', oe_brand='HONDA', part_number='C3', group_id='P3',
        )
        PartCrossCode.objects.create(
            brand='NIBK', product_no='P4', oe_brand='HONDA', part_number='C4', group_id='P4',
        )

    def test_list_shows_unique_brand_names_with_row_counts(self):
        from inventory.services.brand_names_service import BrandNamesServiceCrossCode
        self.assertEqual(BrandNamesServiceCrossCode.list_rows(), [
            {'name': 'ATE', 'row_count': 1},
            {'name': 'HONDA', 'row_count': 3},
            {'name': 'NIBK', 'row_count': 3},
            {'name': 'TOYOTA', 'row_count': 1},
        ])

    def test_search_matches_unique_brand_names(self):
        from inventory.services.brand_names_service import BrandNamesServiceCrossCode
        self.assertEqual(BrandNamesServiceCrossCode.list_rows('honda'), [
            {'name': 'HONDA', 'row_count': 3},
        ])

    def test_search_wildcard_matches_brand_names(self):
        from inventory.services.brand_names_service import BrandNamesServiceCrossCode
        self.assertEqual(BrandNamesServiceCrossCode.list_rows('HON*'), [
            {'name': 'HONDA', 'row_count': 3},
        ])
        self.assertEqual(
            {row['name'] for row in BrandNamesServiceCrossCode.list_rows('N*')},
            {'NIBK'},
        )

    def test_suggestions_prefix_match_brand_names(self):
        from inventory.services.brand_names_service import BrandNamesServiceCrossCode
        self.assertEqual(BrandNamesServiceCrossCode.suggest_names('HON'), ['HONDA'])
        self.assertEqual(BrandNamesServiceCrossCode.suggest_names('NIB'), ['NIBK'])

    def test_rename_updates_every_matching_row(self):
        from inventory.services.brand_names_service import BrandNamesServiceCrossCode
        ok, error, updated_count = BrandNamesServiceCrossCode.rename_brand('HONDA', 'HONDA MOTOR')
        self.assertTrue(ok)
        self.assertEqual(error, '')
        self.assertEqual(updated_count, 3)
        self.assertEqual(PartCrossCode.objects.filter(oe_brand='HONDA MOTOR').count(), 3)
        self.assertEqual(PartCrossCode.objects.filter(oe_brand='HONDA').count(), 0)
        self.assertEqual(PartCrossCode.objects.filter(brand='NIBK').count(), 3)
        self.assertEqual(PartCrossCode.objects.filter(brand='ATE').count(), 1)

    def test_rename_product_brand_updates_every_matching_row(self):
        from inventory.services.brand_names_service import BrandNamesServiceCrossCode
        ok, error, updated_count = BrandNamesServiceCrossCode.rename_brand('NIBK', 'NIBK NEW')
        self.assertTrue(ok)
        self.assertEqual(error, '')
        self.assertEqual(updated_count, 3)
        self.assertEqual(PartCrossCode.objects.filter(brand='NIBK NEW').count(), 3)
        self.assertEqual(PartCrossCode.objects.filter(brand='NIBK').count(), 0)
        self.assertEqual(PartCrossCode.objects.filter(oe_brand='HONDA').count(), 3)

    def test_rename_updates_both_columns_when_name_appears_in_each(self):
        from inventory.services.brand_names_service import BrandNamesServiceCrossCode
        PartCrossCode.objects.create(
            brand='HONDA', product_no='P5', oe_brand='MAZDA', part_number='C5', group_id='P5',
        )
        ok, error, updated_count = BrandNamesServiceCrossCode.rename_brand('HONDA', 'HONDA MOTOR')
        self.assertTrue(ok)
        self.assertEqual(error, '')
        self.assertEqual(updated_count, 4)
        self.assertEqual(PartCrossCode.objects.filter(brand='HONDA').count(), 0)
        self.assertEqual(PartCrossCode.objects.filter(oe_brand='HONDA').count(), 0)
        self.assertEqual(PartCrossCode.objects.filter(brand='HONDA MOTOR').count(), 1)
        self.assertEqual(PartCrossCode.objects.filter(oe_brand='HONDA MOTOR').count(), 3)


class ImportFileHistoryTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(email='hist@example.com', password='Test@123')
        self.other = User.objects.create_user(email='other@example.com', password='Test@123')
        self.client = Client()
        self.client.force_login(self.user)

        self.cars_batch = ImportBatch.objects.create(
            uploaded_by=self.user,
            original_file_name='cars.xlsx',
            import_type=ImportBatch.TYPE_CARS,
            status=ImportBatch.STATUS_COMPLETED,
            total_rows=10,
            cars_count=8,
        )
        self.cc_batch = ImportBatch.objects.create(
            uploaded_by=self.user,
            original_file_name='crosscode.xlsx',
            import_type=ImportBatch.TYPE_PARTS_CROSSCODE,
            status=ImportBatch.STATUS_COMPLETED,
            total_rows=100,
            parts_count=95,
        )
        ImportBatch.objects.create(
            uploaded_by=self.other,
            original_file_name='secret.xlsx',
            import_type=ImportBatch.TYPE_PARTS_CROSSCODE,
            status=ImportBatch.STATUS_COMPLETED,
            total_rows=5,
            parts_count=5,
        )

    def test_batch_helpers(self):
        self.assertEqual(self.cc_batch.imported_records, 95)
        self.assertEqual(self.cc_batch.duplicate_skipped, 5)
        self.assertEqual(self.cc_batch.section_label, 'Cross Code references')
        self.assertEqual(self.cars_batch.section_label, 'Car models')
        self.assertEqual(self.cars_batch.duplicate_skipped, 0)
        self.assertFalse(self.cc_batch.is_deleting)

        self.cc_batch.status = ImportBatch.STATUS_PROCESSING
        self.cc_batch.progress_note = 'Deleting Cross Code parts… (20,000 removed)'
        self.assertTrue(self.cc_batch.is_deleting)
        self.assertEqual(self.cc_batch.display_status, 'Deleting')

    def test_import_page_shows_catalog_history_table(self):
        resp = self.client.get(reverse('inventory:import_data') + '?catalog=crosscode')
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'File import history')
        self.assertContains(resp, 'crosscode.xlsx')
        self.assertContains(resp, 'Cross Code references')
        self.assertContains(resp, '5 duplicate skipped')
        self.assertNotContains(resp, 'cars.xlsx')
        self.assertNotContains(resp, 'secret.xlsx')

        resp_cars = self.client.get(reverse('inventory:import_data') + '?catalog=cars')
        self.assertEqual(resp_cars.status_code, 200)
        self.assertContains(resp_cars, 'cars.xlsx')
        self.assertContains(resp_cars, 'Car models')
        self.assertNotContains(resp_cars, 'crosscode.xlsx')

    def test_history_json_refresh(self):
        resp = self.client.get(reverse('inventory:import_history') + '?catalog=crosscode')
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertIn('table_html', data)
        self.assertIn('crosscode.xlsx', data['table_html'])
        self.assertNotIn('cars.xlsx', data['table_html'])
        self.assertFalse(data['poll'])


class ImportBatchChunkedDeleteTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(email='del@example.com', password='Test@123')
        self.client = Client()
        self.client.force_login(self.user)
        self.batch = ImportBatch.objects.create(
            uploaded_by=self.user,
            original_file_name='big-crosscode.xlsx',
            import_type=ImportBatch.TYPE_PARTS_CROSSCODE,
            status=ImportBatch.STATUS_COMPLETED,
            total_rows=5,
            parts_count=5,
        )
        PartCrossCode.objects.bulk_create([
            PartCrossCode(
                brand='NIBK',
                product_no=f'P{i}',
                oe_brand='MIT',
                part_number=f'CODE{i}',
                group_id=f'P{i}',
                import_batch=self.batch,
            )
            for i in range(5)
        ])
        # Unrelated row must survive delete.
        PartCrossCode.objects.create(
            brand='KEEP',
            product_no='KEEP1',
            oe_brand='KEEP',
            part_number='KEEPCODE',
            group_id='KEEP1',
        )

    def test_chunked_execute_removes_only_batch_rows(self):
        ImportBatchDeleteService.CHUNK_SIZE = 2  # force multiple chunks
        counts = ImportBatchDeleteService.execute(self.batch.id)
        self.assertEqual(counts['parts'], 5)
        self.assertFalse(ImportBatch.objects.filter(pk=self.batch.id).exists())
        self.assertEqual(PartCrossCode.objects.filter(part_number__startswith='CODE').count(), 0)
        self.assertTrue(PartCrossCode.objects.filter(part_number='KEEPCODE').exists())

    def test_delete_view_enqueues_background_delete(self):
        url = reverse('inventory:import_batch_delete', args=[self.batch.id])
        with patch.object(ImportBatchDeleteService, 'enqueue') as mock_enqueue:
            resp = self.client.post(url, {'catalog': 'crosscode'})
        self.assertEqual(resp.status_code, 302)
        mock_enqueue.assert_called_once_with(self.batch.id)

    def test_enqueue_marks_processing_then_execute_clears(self):
        ImportBatchDeleteService.CHUNK_SIZE = 2
        with patch('inventory.services.import_services.threading.Thread') as mock_thread:
            ImportBatchDeleteService.enqueue(self.batch.id)
            mock_thread.assert_called_once()
            # Do not start background thread; run execute ourselves.
        self.batch.refresh_from_db()
        self.assertEqual(self.batch.status, ImportBatch.STATUS_PROCESSING)
        self.assertTrue(self.batch.is_deleting)

        counts = ImportBatchDeleteService.execute(self.batch.id)
        self.assertEqual(counts['parts'], 5)
        self.assertFalse(ImportBatch.objects.filter(pk=self.batch.id).exists())

    def test_chunked_delete_parts_catalog_batch(self):
        batch = ImportBatch.objects.create(
            uploaded_by=self.user,
            original_file_name='parts.xlsx',
            import_type=ImportBatch.TYPE_PARTS,
            status=ImportBatch.STATUS_COMPLETED,
            total_rows=3,
            parts_count=3,
        )
        Part.objects.bulk_create([
            Part(group_id='G1', brand='B', part_number=f'PN{i}', import_batch=batch)
            for i in range(3)
        ])
        ImportBatchDeleteService.CHUNK_SIZE = 1
        counts = ImportBatchDeleteService.execute(batch.id)
        self.assertEqual(counts['parts'], 3)
        self.assertFalse(ImportBatch.objects.filter(pk=batch.id).exists())
        self.assertEqual(Part.objects.filter(part_number__startswith='PN').count(), 0)


class CrossCodeImportUpsertTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(email='cc-upsert@example.com', password='pass')
        self.old_batch = ImportBatch.objects.create(
            uploaded_by=self.user,
            original_file_name='old.xlsx',
            import_type=ImportBatch.TYPE_PARTS_CROSSCODE,
            status=ImportBatch.STATUS_COMPLETED,
        )
        self.new_batch = ImportBatch.objects.create(
            uploaded_by=self.user,
            original_file_name='new.xlsx',
            import_type=ImportBatch.TYPE_PARTS_CROSSCODE,
            status=ImportBatch.STATUS_PROCESSING,
        )
        PartCrossCode.objects.create(
            group_id='PN1',
            brand='NIBK',
            product_no='PN1',
            oe_brand='MITSUBISHI',
            part_number='CODE1',
            import_batch=self.old_batch,
        )

    def test_existing_row_is_skipped_and_only_new_rows_are_saved(self):
        service = ExcelImportService()
        rows = [
            {
                'brand': 'NIBK',
                'product_no': 'PN1',
                'oe_brand': 'MITSUBISHI',
                'part_number': 'CODE1',
                'group_id': 'PN1',
            },
            {
                'brand': 'NIBK',
                'product_no': 'PN2',
                'oe_brand': 'TOYOTA',
                'part_number': 'CODE2',
                'group_id': 'PN2',
            },
        ]
        saved = service._insert_new_crosscode_parts_batched(rows, self.new_batch, bs=100)
        self.assertEqual(saved, 1)
        self.assertEqual(PartCrossCode.objects.count(), 2)

        existing = PartCrossCode.objects.get(product_no='PN1', part_number='CODE1')
        self.assertEqual(existing.import_batch_id, self.old_batch.id)

        created = PartCrossCode.objects.get(product_no='PN2', part_number='CODE2')
        self.assertEqual(created.import_batch_id, self.new_batch.id)
        self.assertEqual(created.oe_brand, 'TOYOTA')

    def test_all_duplicates_save_zero_and_history_counts_skipped(self):
        service = ExcelImportService()
        rows = [
            {
                'brand': 'NIBK',
                'product_no': 'PN1',
                'oe_brand': 'MITSUBISHI',
                'part_number': 'CODE1',
                'group_id': 'PN1',
            },
        ]
        saved = service._insert_new_crosscode_parts_batched(rows, self.new_batch, bs=100)
        self.assertEqual(saved, 0)
        self.assertEqual(PartCrossCode.objects.count(), 1)
        existing = PartCrossCode.objects.get(product_no='PN1', part_number='CODE1')
        self.assertEqual(existing.import_batch_id, self.old_batch.id)

        self.new_batch.total_rows = 51357
        self.new_batch.parts_count = saved
        self.new_batch.status = ImportBatch.STATUS_COMPLETED
        self.assertEqual(self.new_batch.imported_records, 0)
        self.assertEqual(self.new_batch.duplicate_skipped, 51357)


class CrossCarsSourceLabelTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(email='source@example.com', password='pass')
        self.catalog_batch = ImportBatch.objects.create(
            uploaded_by=self.user,
            original_file_name='parts.xlsx',
            import_type=ImportBatch.TYPE_PARTS,
            status=ImportBatch.STATUS_COMPLETED,
        )
        self.cwp_batch = ImportBatch.objects.create(
            uploaded_by=self.user,
            original_file_name='cwp.xlsx',
            import_type=ImportBatch.TYPE_CAR_WITH_PARTS,
            status=ImportBatch.STATUS_COMPLETED,
        )

        self.car_catalog = Car.objects.create(car_id='C1', car_model='TOYOTA Catalog Car')
        self.car_cwp = Car.objects.create(car_id='C2', car_model='TOYOTA CWP Car')
        self.car_manual = Car.objects.create(car_id='C3', car_model='TOYOTA Manual Car')

        CarGroup.objects.create(car_id=str(self.car_catalog.id), group_id='G-CAT')
        CarGroup.objects.create(car_id=str(self.car_cwp.id), group_id=f'CWP{self.car_cwp.id}')
        CarGroup.objects.create(car_id=str(self.car_manual.id), group_id=f'MANUAL{self.car_manual.id}')

        Part.objects.create(
            group_id='G-CAT', brand='OEM', part_number='SRC001', import_batch=self.catalog_batch,
        )
        Part.objects.create(
            group_id=f'CWP{self.car_cwp.id}', brand='OEM', part_number='SRC001', import_batch=self.cwp_batch,
        )
        Part.objects.create(
            group_id=f'MANUAL{self.car_manual.id}', brand='OEM', part_number='SRC001',
        )

    def test_resolve_part_source_labels(self):
        catalog_part = Part.objects.get(group_id='G-CAT')
        cwp_part = Part.objects.get(group_id=f'CWP{self.car_cwp.id}')
        manual_part = Part.objects.get(group_id=f'MANUAL{self.car_manual.id}')
        self.assertEqual(PartSearchService.resolve_part_source(catalog_part), SOURCE_CATALOG_IMPORT)
        self.assertEqual(PartSearchService.resolve_part_source(cwp_part), SOURCE_CAR_WITH_PARTS)
        self.assertEqual(PartSearchService.resolve_part_source(manual_part), SOURCE_MANUAL_ENTRY)

    def test_search_results_split_by_import_source(self):
        results = PartSearchService.build_results('SRC001')
        by_source = {item['source']: item for item in results}
        self.assertEqual(set(by_source), {
            SOURCE_CATALOG_IMPORT,
            SOURCE_CAR_WITH_PARTS,
            SOURCE_MANUAL_ENTRY,
        })
        self.assertEqual(by_source[SOURCE_CATALOG_IMPORT]['car'].id, self.car_catalog.id)
        self.assertEqual(by_source[SOURCE_CAR_WITH_PARTS]['car'].id, self.car_cwp.id)
        self.assertEqual(by_source[SOURCE_MANUAL_ENTRY]['car'].id, self.car_manual.id)


class CarWithPartsImportFormatTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(email='cwp@example.com', password='pass')
        self.batch = ImportBatch.objects.create(
            uploaded_by=self.user,
            original_file_name='cwp.xlsx',
            import_type=ImportBatch.TYPE_CAR_WITH_PARTS,
            status=ImportBatch.STATUS_PROCESSING,
        )
        self.service = ExcelImportService()

    def test_sample_workbook_has_brand_part_model_columns(self):
        workbook = ExcelImportService.build_car_with_parts_sample_workbook()
        sheet = workbook.active
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        self.assertEqual(headers, ['Brand Name', 'Part Number', 'Model'])
        first_data = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]
        self.assertEqual(len(first_data), 3)
        self.assertTrue(first_data[0])  # brand
        self.assertTrue(first_data[1])  # part
        self.assertTrue(first_data[2])  # model

    def test_parse_brand_part_model_rows(self):
        from inventory.services.spreadsheet_loader import _RowsSheet

        sheet = _RowsSheet([
            ('Brand Name', 'Part Number', 'Model'),
            ('TOYOTA', 'PN100', 'TOYOTA COROLLA AE101'),
            ('HONDA', 'PN200', 'HONDA CIVIC EG'),
            ('', '', ''),
            ('TOYOTA', '', 'TOYOTA COROLLA AE101'),  # missing part → warning/skip
        ])
        parsed = self.service._parse_car_with_parts_sheet(sheet, 'Car with parts')
        self.assertEqual(parsed['total_rows'], 2)
        self.assertEqual(parsed['car_names'], {'TOYOTA COROLLA AE101', 'HONDA CIVIC EG'})
        self.assertEqual(
            {(r['brand'], r['part_number'], r['model']) for r in parsed['rows']},
            {
                ('TOYOTA', 'PN100', 'TOYOTA COROLLA AE101'),
                ('HONDA', 'PN200', 'HONDA CIVIC EG'),
            },
        )
        self.assertTrue(self.service.errors)

    def test_import_creates_car_and_part_with_brand(self):
        rows = [
            {'model': 'TOYOTA COROLLA AE101', 'brand': 'TOYOTA', 'part_number': 'PN100'},
            {'model': 'TOYOTA COROLLA AE101', 'brand': 'TOYOTA', 'part_number': 'PN101'},
        ]
        car_names = {'TOYOTA COROLLA AE101'}
        cars_created, car_lookup = self.service._upsert_car_with_parts_cars(car_names, self.batch, bs=100)
        self.assertEqual(cars_created, 1)
        group_by_car_name = {}
        self.service._upsert_car_with_parts_groups(car_lookup, group_by_car_name, self.batch, bs=100)
        parts_created = self.service._upsert_car_with_parts_parts(rows, group_by_car_name, self.batch, bs=100)
        self.assertEqual(parts_created, 2)

        car = Car.objects.get(car_id='TOYOTA COROLLA AE101')
        self.assertEqual(car.car_model, 'TOYOTA COROLLA AE101')
        parts = list(Part.objects.filter(group_id=f'CWP{car.id}').order_by('part_number'))
        self.assertEqual(len(parts), 2)
        self.assertEqual(parts[0].brand, 'TOYOTA')
        self.assertEqual(parts[0].part_number, 'PN100')
        self.assertEqual(parts[1].brand, 'TOYOTA')
        self.assertEqual(parts[1].part_number, 'PN101')

    def test_reimport_updates_existing_part_brand(self):
        car = Car.objects.create(car_id='HONDA CIVIC EG', car_model='HONDA CIVIC EG')
        group_id = f'CWP{car.id}'
        CarGroup.objects.create(car_id=str(car.id), group_id=group_id)
        Part.objects.create(group_id=group_id, brand='', part_number='PN200')

        rows = [{'model': 'HONDA CIVIC EG', 'brand': 'HONDA', 'part_number': 'PN200'}]
        car_lookup = {'HONDA CIVIC EG': car.id}
        group_by_car_name = {'HONDA CIVIC EG': group_id}
        created = self.service._upsert_car_with_parts_parts(rows, group_by_car_name, self.batch, bs=100)
        self.assertEqual(created, 0)
        part = Part.objects.get(group_id=group_id, part_number='PN200')
        self.assertEqual(part.brand, 'HONDA')
        self.assertEqual(part.import_batch_id, self.batch.id)


class CarCatalogKeywordSearchTests(TestCase):
    def setUp(self):
        Car.objects.create(
            car_id='T1',
            car_model='TOYOTA LAND CRUISER UZZ100 1998-2007',
            transmission='AT',
            engine='4.7',
        )
        Car.objects.create(
            car_id='T2',
            car_model='TOYOTA COROLLA AE101 1991-2000',
            transmission='MT',
            engine='1.6',
        )
        Car.objects.create(
            car_id='H1',
            car_model='HONDA CIVIC EG 1991-1995',
            transmission='MT',
            engine='1.5',
        )

    def test_non_contiguous_keywords_match_model(self):
        from inventory.services.car_catalog_service import CarCatalogService

        qs = CarCatalogService.build_queryset({'q': 'TOYOTA UZZ', 'brand': '', 'model': '', 'engine': '', 'transmission': '', 'steering': '', 'wd': '', 'parameters': ''})
        models = list(qs.values_list('car_model', flat=True))
        self.assertEqual(len(models), 1)
        self.assertIn('UZZ100', models[0])
        self.assertNotIn('COROLLA', models[0])

    def test_suggestions_use_same_keyword_rules(self):
        from inventory.services.car_catalog_service import CarCatalogService

        suggestions = CarCatalogService.get_suggestions({}, 'TOYOTA UZZ')
        self.assertEqual(len(suggestions), 1)
        self.assertIn('UZZ100', suggestions[0])

        suggestions_brand = CarCatalogService.get_suggestions({'brand': 'TOYOTA'}, 'COROLLA')
        self.assertEqual(len(suggestions_brand), 1)
        self.assertIn('COROLLA', suggestions_brand[0])

    def test_manual_add_suggestions_match_same_rules(self):
        suggestions = ManualEntryService.suggest_car_models('TOYOTA UZZ')
        self.assertEqual(len(suggestions), 1)
        self.assertIn('UZZ100', suggestions[0])


class CrossCarsBulkSearchFormatTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(email='bulk@example.com', password='pass')
        self.car = Car.objects.create(car_id='C-BULK', car_model='TOYOTA TEST CAR')
        CarGroup.objects.create(car_id=str(self.car.id), group_id='G-BULK')
        Part.objects.create(group_id='G-BULK', brand='TOYOTA', part_number='3231A047')

    def test_sample_workbook_columns(self):
        workbook = BulkSearchService.build_sample_workbook()
        headers = [cell.value for cell in next(workbook.active.iter_rows(min_row=1, max_row=1))]
        self.assertEqual(headers, ['Brand Name', 'Brand Number', 'Part Number'])

    def test_parse_header_aliases_brand_name_brand_number(self):
        brand_idx, brand_number_idx, part_number_idx = BulkSearchService._parse_header_indices(
            ['Brand Name', 'Brand Number', 'Part Number']
        )
        self.assertEqual(brand_idx, 0)
        self.assertEqual(brand_number_idx, 1)
        self.assertEqual(part_number_idx, 2)

    def test_parse_upload_and_run_bulk_search(self):
        import openpyxl
        from io import BytesIO

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(['Brand Name', 'Brand Number', 'Part Number'])
        sheet.append(['AISIN', 'AS-12345', '3231A047'])
        sheet.append(['AISIN', 'AS-99999', 'NOTFOUND999'])
        buffer = BytesIO()
        workbook.save(buffer)
        upload = SimpleUploadedFile(
            'bulk.xlsx',
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

        rows_data, error = BulkSearchService.parse_upload(upload)
        self.assertIsNone(error)
        self.assertEqual(len(rows_data), 2)
        self.assertEqual(rows_data[0]['brand'], 'AISIN')
        self.assertEqual(rows_data[0]['brand_number'], 'AS-12345')
        self.assertEqual(rows_data[0]['part_number'], '3231A047')

        results, summary, err = BulkSearchService.run_bulk_search(
            self.user, rows_data, save_to_basket=False,
        )
        self.assertIsNone(err)
        statuses = {r['part_number']: r['status'] for r in results}
        self.assertEqual(statuses['3231A047'], 'Found')
        self.assertEqual(statuses['NOTFOUND999'], 'Not Found')
        self.assertEqual(summary['found_count'], 1)
        self.assertEqual(summary['searched_count'], 2)

    def test_missing_columns_returns_error(self):
        import openpyxl
        from io import BytesIO

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(['Part Number'])
        sheet.append(['3231A047'])
        buffer = BytesIO()
        workbook.save(buffer)
        upload = SimpleUploadedFile(
            'bad.xlsx',
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        rows_data, error = BulkSearchService.parse_upload(upload)
        self.assertIsNone(rows_data)
        self.assertIn('Brand Name', error)


class CrossCodeBasketFamilyListTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(email='basket@example.com', password='Test@123')
        self.client.force_login(self.user)
        self.part_a = PartCrossCode.objects.create(
            brand='NIBK', product_no='P230', oe_brand='MIT', part_number='CODE-A', group_id='P230',
        )
        self.part_b = PartCrossCode.objects.create(
            brand='NIBK', product_no='P230', oe_brand='TOYOTA', part_number='CODE-B', group_id='P230',
        )
        BasketServiceCrossCode.add_item(self.user, self.part_a, 'KANOYA', '230001')
        BasketServiceCrossCode.add_item(self.user, self.part_b, 'KANOYA', '230001')

    def test_basket_lists_family_rows_without_open_group_table(self):
        response = self.client.get(reverse('inventory:basket_crosscode'))
        self.assertContains(response, 'id="basket-table"')
        self.assertContains(response, '<th>Product Brand</th>')
        self.assertContains(response, '<th>Product No</th>')
        self.assertContains(response, '<th>Brand</th>')
        self.assertContains(response, '<th>Code</th>')
        self.assertContains(response, 'CODE-A')
        self.assertContains(response, 'CODE-B')
        self.assertNotContains(response, '<th>Cross Brand</th>')
        self.assertNotContains(response, '<th>Cross Code</th>')
        self.assertNotContains(response, 'btn-view-group')
        self.assertNotContains(response, '>Open<')

    def test_basket_search_by_cross_code_shows_family_rows(self):
        response = self.client.get(reverse('inventory:basket_crosscode'), {'q': '230001'})
        self.assertContains(response, 'CODE-A')
        self.assertContains(response, 'CODE-B')
        self.assertContains(response, 'matches found')

    def test_basket_search_by_brand_code_finds_new_row(self):
        response = self.client.get(reverse('inventory:basket_crosscode'), {'q': 'CODE-B'})
        self.assertContains(response, 'CODE-B')
        self.assertNotContains(response, 'CODE-A')

    def test_basket_search_by_entered_brand_finds_new_rows(self):
        response = self.client.get(reverse('inventory:basket_crosscode'), {'q': 'KANOYA'})
        self.assertContains(response, 'CODE-A')
        self.assertContains(response, 'CODE-B')

    def test_basket_wildcard_search_matches_entered_code(self):
        response = self.client.get(reverse('inventory:basket_crosscode'), {'q': '230*'})
        self.assertContains(response, 'CODE-A')
        self.assertContains(response, 'CODE-B')
        self.assertContains(response, 'matches found')

    def test_basket_wildcard_search_matches_catalog_code(self):
        response = self.client.get(reverse('inventory:basket_crosscode'), {'q': 'CODE-*'})
        self.assertContains(response, 'CODE-A')
        self.assertContains(response, 'CODE-B')

    def test_basket_wildcard_search_matches_entered_brand(self):
        response = self.client.get(reverse('inventory:basket_crosscode'), {'q': 'KAN*'})
        self.assertContains(response, 'CODE-A')
        self.assertContains(response, 'CODE-B')

    def test_group_detail_redirects_to_family_list(self):
        basket = BasketCrossCode.objects.get(brand='KANOYA', brand_number='230001')
        response = self.client.get(
            reverse('inventory:basket_group_detail_crosscode', kwargs={'basket_id': basket.pk}),
        )
        self.assertRedirects(response, reverse('inventory:basket_crosscode'))

    def test_basket_suggestions_prefix_match_entered_codes(self):
        BasketServiceCrossCode.add_item(self.user, self.part_a, 'OTHER', '230002')
        suggestions = BasketServiceCrossCode.suggest_values(self.user, '23000')
        self.assertEqual(suggestions, ['230001', '230002'])

    def test_basket_suggestions_endpoint_returns_prefix_matches(self):
        response = self.client.get(
            reverse('inventory:basket_suggestions_crosscode'),
            {'q': '23000'},
        )
        self.assertEqual(response.json()['suggestions'], ['230001'])

