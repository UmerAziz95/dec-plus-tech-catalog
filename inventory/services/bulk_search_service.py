from django.conf import settings
import openpyxl
import tempfile
from pathlib import Path

from inventory.models import (
    BasketItem, BasketItemCrossCode, Car, CarCrossCode, CarGroup, CarGroupCrossCode,
    Part, PartCrossCode,
)
from inventory.services.basket_service import BasketService, BasketServiceCrossCode
from inventory.services.part_number_utils import sanitize_crosscode_bulk_part, sanitize_part_number
from inventory.services.spreadsheet_loader import SUPPORTED_EXTENSIONS, load_workbook


class BulkSearchService:
    car_model = Car
    car_group_model = CarGroup
    part_model = Part
    basket_item_model = BasketItem
    basket_service = BasketService

    @staticmethod
    def _parse_header_indices(header_row):
        brand_idx = -1
        brand_number_idx = -1
        part_number_idx = -1

        for i, header in enumerate(header_row):
            if not header:
                continue
            h = str(header).lower().strip()
            if h in ['brand', 'cross brand']:
                brand_idx = i
            elif h in ['brand number', 'cross code', 'code']:
                brand_number_idx = i
            elif h in ['part number', 'part_number']:
                part_number_idx = i

        return brand_idx, brand_number_idx, part_number_idx

    @classmethod
    def _read_rows(cls, sheet, brand_idx, brand_number_idx, part_number_idx):
        rows_data = []
        part_numbers_to_search = set()

        for row in sheet.iter_rows(min_row=2, values_only=True):
            part_no = row[part_number_idx]
            if not part_no:
                continue

            part_no = str(part_no).strip()
            part_no = sanitize_part_number(part_no)
            if not part_no:
                continue
            brand = str(row[brand_idx]).strip() if brand_idx != -1 and row[brand_idx] else ''
            brand_num = str(row[brand_number_idx]).strip() if brand_number_idx != -1 and row[brand_number_idx] else ''
            brand, brand_num = cls.basket_service.normalize_cross_brand_fields(brand, brand_num)

            rows_data.append({
                'brand': brand,
                'brand_number': brand_num,
                'part_number': part_no,
            })
            part_numbers_to_search.add(part_no)

        return rows_data, part_numbers_to_search

    @classmethod
    def _load_matches(cls, part_numbers_to_search):
        """
        Exact part-number match only for bulk search.

        Normalized equality is required, so searching "54500-8H31"
        (norm: 545008H31) will NOT match "54500-8H31A" / "54500-8H31B".
        Manual search keeps its own contains fallback and is unaffected.
        """
        if not part_numbers_to_search:
            return {}, {}

        # Use the stored generated column part_number_norm with B-tree index.
        # psycopg3 doesn't auto-expand a tuple param into "(a, b, c)" for
        # "IN %s" the way psycopg2 did, so build explicit placeholders.
        search_keys = {sanitize_part_number(p) for p in part_numbers_to_search if p}
        search_keys.discard('')
        if not search_keys:
            return {}, {}

        placeholders = ', '.join(['%s'] * len(search_keys))
        matching_parts = cls.part_model.objects.extra(
            where=[f"part_number_norm IN ({placeholders})"],
            params=list(search_keys),
        ).order_by('part_number', 'group_id').distinct('part_number', 'group_id')

        parts_by_part_num = {}
        group_pks = set()
        for part in matching_parts:
            # Defensive: only accept exact normalized equality.
            pn = sanitize_part_number(part.part_number)
            if pn not in search_keys:
                continue
            if pn not in parts_by_part_num:
                parts_by_part_num[pn] = []
            parts_by_part_num[pn].append(part)
            group_pks.add(part.group_id)

        car_groups = cls.car_group_model.objects.filter(group_id__in=group_pks)

        # CarGroup.car_id actually stores the Car primary key (as text),
        # not the business Car.car_id string — filter/lookup by pk.
        car_pks = set(car_groups.values_list('car_id', flat=True))
        cars = cls.car_model.objects.filter(id__in=car_pks)
        cars_by_car_id = {str(c.id): c for c in cars}

        cars_by_group_pk = {}
        for cg in car_groups:
            car = cars_by_car_id.get(cg.car_id)
            if not car:
                continue
            if cg.group_id not in cars_by_group_pk:
                cars_by_group_pk[cg.group_id] = []
            cars_by_group_pk[cg.group_id].append(car)

        return parts_by_part_num, cars_by_group_pk

    @classmethod
    def _existing_basket_keys(cls, user, part_numbers_to_search):
        return cls.basket_service.existing_item_keys(user, part_numbers_to_search)

    @classmethod
    def _build_basket_entries(cls, user, rows_data, parts_by_part_num, cars_by_group_pk, existing_keys):
        entries = []
        pending_keys = set()
        batch_size = getattr(settings, 'IMPORT_ROW_BATCH_SIZE', 200)
        brand_pairs = {
            cls.basket_service.normalize_cross_brand_fields(
                row_data['brand'],
                row_data['brand_number'],
            )
            for row_data in rows_data
            if row_data['brand'] and row_data['brand_number']
        }
        basket_map = cls.basket_service.get_or_create_baskets_for_pairs(brand_pairs)

        for row_data in rows_data:
            brand, brand_number = cls.basket_service.normalize_cross_brand_fields(
                row_data['brand'],
                row_data['brand_number'],
            )
            part_no = row_data['part_number']

            if not brand or not brand_number:
                continue

            found_parts = parts_by_part_num.get(part_no, [])
            if not found_parts:
                continue

            basket = basket_map[(brand, brand_number)]
            for part in found_parts:
                for car in cars_by_group_pk.get(part.group_id, []):
                    key = (car.id, part.part_number, basket.id)
                    if key in existing_keys or key in pending_keys:
                        continue
                    pending_keys.add(key)
                    entries.append(cls.basket_item_model(
                        user=user,
                        car=car,
                        part=part,
                        basket=basket,
                        group_id=part.group_id,
                    ))

        for offset in range(0, len(entries), batch_size):
            cls.basket_item_model.objects.bulk_create(entries[offset:offset + batch_size], ignore_conflicts=True)

        return len(entries)

    @staticmethod
    def _matched_car_ids(rows_data, parts_by_part_num, cars_by_group_pk):
        matched_car_ids = set()
        for row_data in rows_data:
            if not row_data['brand'] or not row_data['brand_number']:
                continue
            for part in parts_by_part_num.get(row_data['part_number'], []):
                for car in cars_by_group_pk.get(part.group_id, []):
                    matched_car_ids.add(car.id)
        return matched_car_ids

    @staticmethod
    def _build_results(rows_data, parts_by_part_num, cars_by_group_pk, basket_part_numbers):
        results = []

        for row_data in rows_data:
            part_no = row_data['part_number']
            brand = row_data['brand']
            brand_number = row_data['brand_number']
            found_parts = parts_by_part_num.get(part_no, [])

            if not brand or not brand_number:
                results.append({
                    'brand': brand,
                    'brand_number': brand_number,
                    'part_number': part_no,
                    'car_count': 0,
                    'top_cars': [],
                    'in_basket': False,
                    'status': 'Not Found',
                })
                continue

            if not found_parts:
                results.append({
                    'brand': brand,
                    'brand_number': brand_number,
                    'part_number': part_no,
                    'car_count': 0,
                    'top_cars': [],
                    'in_basket': False,
                    'status': 'Not Found',
                })
                continue

            car_set = set()
            top_cars = []
            for part in found_parts:
                for car in cars_by_group_pk.get(part.group_id, []):
                    if car.id not in car_set:
                        car_set.add(car.id)
                        if len(top_cars) < 3:
                            top_cars.append(car.car_model)

            results.append({
                'brand': brand,
                'brand_number': brand_number,
                'part_number': part_no,
                'car_count': len(car_set),
                'top_cars': top_cars,
                'in_basket': part_no in basket_part_numbers,
                'status': 'Found',
            })

        return results

    @staticmethod
    def _build_summary(rows_data, parts_by_part_num, matched_car_ids, missed_rows, added_to_basket):
        searched_count = len(rows_data)
        found_count = 0
        for row_data in rows_data:
            if row_data['brand'] and row_data['brand_number'] and row_data['part_number'] in parts_by_part_num:
                found_count += 1

        return {
            'searched_count': searched_count,
            'found_count': found_count,
            'car_models_matched': len(matched_car_ids),
            'missed_count': len(missed_rows),
            'added_to_basket': added_to_basket,
            'missed_rows': missed_rows,
        }

    @classmethod
    def parse_upload(cls, excel_file):
        name = (excel_file.name or '').lower()
        if not name.endswith(SUPPORTED_EXTENSIONS):
            allowed = ', '.join(SUPPORTED_EXTENSIONS)
            return None, f'Please upload a valid file ({allowed}).'

        tmp_path = None
        try:
            suffix = Path(name).suffix.lower() or '.xlsx'
            with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
                for chunk in excel_file.chunks():
                    tmp.write(chunk)
                tmp_path = tmp.name

            workbook = load_workbook(tmp_path)
            sheet = workbook[workbook.sheetnames[0]]
            header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            brand_idx, brand_number_idx, part_number_idx = cls._parse_header_indices(header_row)

            if part_number_idx == -1:
                return None, 'Could not find "Part Number" or "Part_number" column in the file.'

            rows_data, _part_numbers = cls._read_rows(sheet, brand_idx, brand_number_idx, part_number_idx)
            return rows_data, None
        except Exception as exc:
            return None, f'Error processing file: {str(exc)}'
        finally:
            if tmp_path:
                try:
                    Path(tmp_path).unlink(missing_ok=True)
                except OSError:
                    pass

    @classmethod
    def run_bulk_search(cls, user, rows_data, save_to_basket=True):
        if not rows_data:
            summary = {
                'searched_count': 0,
                'found_count': 0,
                'car_models_matched': 0,
                'missed_count': 0,
                'added_to_basket': 0,
                'missed_rows': [],
            }
            return [], summary, None

        part_numbers_to_search = {row['part_number'] for row in rows_data if row.get('part_number')}
        parts_by_part_num, cars_by_group_pk = cls._load_matches(part_numbers_to_search)

        added_to_basket = 0
        basket_part_numbers = set()
        if save_to_basket:
            existing_keys = cls._existing_basket_keys(user, part_numbers_to_search)
            added_to_basket = cls._build_basket_entries(
                user,
                rows_data,
                parts_by_part_num,
                cars_by_group_pk,
                existing_keys,
            )
            placeholders = ', '.join(['%s'] * len(part_numbers_to_search))
            matched_part_ids = cls.part_model.objects.extra(
                where=[f"part_number_norm IN ({placeholders})"],
                params=list(part_numbers_to_search),
            ).values_list('id', flat=True)
            basket_part_numbers = {
                sanitize_part_number(part_number)
                for part_number in cls.basket_item_model.objects.filter(
                    user=user,
                    part_id__in=matched_part_ids,
                ).values_list('part__part_number', flat=True)
            }

        matched_car_ids = cls._matched_car_ids(rows_data, parts_by_part_num, cars_by_group_pk)
        results = cls._build_results(rows_data, parts_by_part_num, cars_by_group_pk, basket_part_numbers)

        missed_rows = []
        for row_data in rows_data:
            part_no = row_data['part_number']
            brand = row_data['brand']
            brand_number = row_data['brand_number']
            if not brand or not brand_number or part_no not in parts_by_part_num:
                missed_rows.append(row_data)

        summary = cls._build_summary(rows_data, parts_by_part_num, matched_car_ids, missed_rows, added_to_basket)
        return results, summary, None

    @classmethod
    def process_upload(cls, user, excel_file):
        rows_data, error_message = cls.parse_upload(excel_file)
        if error_message:
            return None, False, None, error_message
        results, summary, err = cls.run_bulk_search(user, rows_data, save_to_basket=True)
        return results, True, summary, err

    @staticmethod
    def build_sample_workbook():
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = 'Bulk search sample'
        sheet.append(['Cross Brand', 'Cross Code', 'Part Number'])
        sheet.append(['AISIN', 'AS-12345', '3231A047'])
        return workbook

    @staticmethod
    def build_missed_workbook(missed_rows):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = 'Missed'
        sheet.append(['Cross Brand', 'Cross Code', 'Part Number'])
        for row in missed_rows:
            sheet.append([
                row.get('brand', ''),
                row.get('brand_number', ''),
                row.get('part_number', ''),
            ])
        return workbook

    @staticmethod
    def build_results_workbook(results):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = 'Bulk search'
        sheet.append([
            'Cross Brand',
            'Cross Code',
            'Part number',
            'Status',
            'Vehicle count',
            'Sample models',
            'In basket',
        ])
        for row in results:
            top = row.get('top_cars') or []
            sample = '; '.join(top[:10])
            if len(top) > 10:
                sample = f'{sample}...'
            sheet.append([
                str(row.get('brand') or ''),
                str(row.get('brand_number') or ''),
                str(row.get('part_number') or ''),
                str(row.get('status') or ''),
                int(row.get('car_count') or 0),
                sample,
                'Yes' if row.get('in_basket') else 'No',
            ])
        return workbook


class BulkSearchServiceCrossCode(BulkSearchService):
    car_model = CarCrossCode
    car_group_model = CarGroupCrossCode
    part_model = PartCrossCode
    basket_item_model = BasketItemCrossCode
    basket_service = BasketServiceCrossCode

    @classmethod
    def _read_rows(cls, sheet, brand_idx, brand_number_idx, part_number_idx):
        """
        Cross Code bulk rules:
        - Strip symbols/spaces; star is removed from the cleaned value.
        - Any row that contained "*" is auto-missed (no search).
        - Remaining rows are exact-match only.
        """
        rows_data = []
        part_numbers_to_search = set()

        for row in sheet.iter_rows(min_row=2, values_only=True):
            part_no = row[part_number_idx]
            if not part_no:
                continue

            raw = str(part_no).strip()
            if not raw:
                continue

            brand = str(row[brand_idx]).strip() if brand_idx != -1 and row[brand_idx] else ''
            brand_num = (
                str(row[brand_number_idx]).strip()
                if brand_number_idx != -1 and row[brand_number_idx]
                else ''
            )
            brand, brand_num = cls.basket_service.normalize_cross_brand_fields(brand, brand_num)

            cleaned, had_star = sanitize_crosscode_bulk_part(raw)
            if had_star or not cleaned:
                rows_data.append({
                    'brand': brand,
                    'brand_number': brand_num,
                    'part_number': raw if had_star else cleaned,
                    'bulk_miss': True,
                })
                continue

            rows_data.append({
                'brand': brand,
                'brand_number': brand_num,
                'part_number': cleaned,
                'bulk_miss': False,
            })
            part_numbers_to_search.add(cleaned)

        return rows_data, part_numbers_to_search

    @classmethod
    def run_bulk_search(cls, user, rows_data, save_to_basket=True):
        if not rows_data:
            summary = {
                'searched_count': 0,
                'found_count': 0,
                'car_models_matched': 0,
                'missed_count': 0,
                'added_to_basket': 0,
                'missed_rows': [],
            }
            return [], summary, None

        searchable_rows = [row for row in rows_data if not row.get('bulk_miss')]
        part_numbers_to_search = {
            row['part_number']
            for row in searchable_rows
            if row.get('part_number')
        }
        part_numbers_to_search.discard('')
        parts_by_part_num, cars_by_group_pk = cls._load_matches(part_numbers_to_search)

        added_to_basket = 0
        basket_part_numbers = set()
        if save_to_basket:
            existing_keys = cls._existing_basket_keys(user, part_numbers_to_search)
            added_to_basket = cls._build_basket_entries(
                user,
                searchable_rows,
                parts_by_part_num,
                cars_by_group_pk,
                existing_keys,
            )
            matched_part_ids = {
                part.id
                for parts in parts_by_part_num.values()
                for part in parts
            }
            basket_part_numbers = {
                sanitize_part_number(part_number)
                for part_number in cls.basket_item_model.objects.filter(
                    user=user,
                    part_id__in=matched_part_ids,
                ).values_list('part__part_number', flat=True)
            }
            for key, parts in parts_by_part_num.items():
                if any(sanitize_part_number(p.part_number) in basket_part_numbers for p in parts):
                    basket_part_numbers.add(key)

        matched_car_ids = cls._matched_car_ids(rows_data, parts_by_part_num, cars_by_group_pk)
        results = cls._build_results(rows_data, parts_by_part_num, cars_by_group_pk, basket_part_numbers)

        missed_rows = []
        for row_data in rows_data:
            if row_data.get('bulk_miss'):
                missed_rows.append(row_data)
                continue
            part_no = row_data['part_number']
            brand = row_data['brand']
            brand_number = row_data['brand_number']
            if not brand or not brand_number or part_no not in parts_by_part_num:
                missed_rows.append(row_data)

        summary = cls._build_summary(
            rows_data, parts_by_part_num, matched_car_ids, missed_rows, added_to_basket,
        )
        return results, summary, None

    @classmethod
    def _load_matches(cls, part_numbers_to_search):
        """Exact match on Code / Product No, then expand each product family."""
        from inventory.services.part_number_utils import filter_parts_exact
        from inventory.services.part_search_service import PartSearchServiceCrossCode

        parts_by_part_num = {}
        search_keys = {sanitize_part_number(p) for p in part_numbers_to_search if p}
        search_keys.discard('')
        if not search_keys:
            return {}, {}

        for key in search_keys:
            found = list(filter_parts_exact(key, model=cls.part_model))
            if not found:
                found = list(
                    cls.part_model.objects.extra(
                        where=[
                            "regexp_replace(upper(coalesce(product_no, '')), %s, '', 'g') = %s"
                        ],
                        params=[r'[^A-Za-z0-9]', key],
                    )[:500]
                )
            if found:
                parts_by_part_num[key] = PartSearchServiceCrossCode.expand_product_families(found)
        return parts_by_part_num, {}

    @classmethod
    def _build_basket_entries(cls, user, rows_data, parts_by_part_num, cars_by_group_pk, existing_keys):
        entries = []
        pending_keys = set()
        batch_size = getattr(settings, 'IMPORT_ROW_BATCH_SIZE', 200)
        brand_pairs = {
            cls.basket_service.normalize_cross_brand_fields(
                row_data['brand'],
                row_data['brand_number'],
            )
            for row_data in rows_data
            if row_data['brand'] and row_data['brand_number'] and not row_data.get('bulk_miss')
        }
        basket_map = cls.basket_service.get_or_create_baskets_for_pairs(brand_pairs)

        # Reload existing keys for every expanded Brand+Code in the families.
        all_part_numbers = {
            part.part_number
            for parts in parts_by_part_num.values()
            for part in parts
            if part.part_number
        }
        if all_part_numbers:
            existing_keys = cls.basket_service.existing_item_keys(user, all_part_numbers)

        for row_data in rows_data:
            if row_data.get('bulk_miss'):
                continue
            brand, brand_number = cls.basket_service.normalize_cross_brand_fields(
                row_data['brand'],
                row_data['brand_number'],
            )
            part_no = sanitize_part_number(row_data['part_number'])
            if not brand or not brand_number:
                continue
            found_parts = parts_by_part_num.get(part_no, [])
            if not found_parts:
                continue
            basket = basket_map[(brand, brand_number)]
            # One basket line per Brand + Code under this Brand Name + Brand Number.
            for part in found_parts:
                key = ((part.oe_brand or ''), part.part_number, basket.id)
                if key in existing_keys or key in pending_keys:
                    continue
                pending_keys.add(key)
                entries.append(cls.basket_item_model(
                    user=user,
                    car=None,
                    part=part,
                    basket=basket,
                    group_id=part.group_id or part.product_no or '',
                ))

        for offset in range(0, len(entries), batch_size):
            cls.basket_item_model.objects.bulk_create(
                entries[offset:offset + batch_size],
                ignore_conflicts=True,
            )
        return len(entries)

    @staticmethod
    def _matched_car_ids(rows_data, parts_by_part_num, cars_by_group_pk):
        return set()

    @staticmethod
    def _build_results(rows_data, parts_by_part_num, cars_by_group_pk, basket_part_numbers):
        results = []
        for row_data in rows_data:
            brand = row_data['brand']
            brand_number = row_data['brand_number']
            if row_data.get('bulk_miss'):
                results.append({
                    'brand': brand,
                    'brand_number': brand_number,
                    'part_number': row_data['part_number'],
                    'status': 'Missed (*)',
                    'car_count': 0,
                    'top_cars': [],
                    'in_basket': False,
                })
                continue
            part_no = sanitize_part_number(row_data['part_number'])
            found_parts = parts_by_part_num.get(part_no, [])
            if not brand or not brand_number:
                results.append({
                    'brand': brand,
                    'brand_number': brand_number,
                    'part_number': part_no,
                    'status': 'Skipped',
                    'car_count': 0,
                    'top_cars': [],
                    'in_basket': False,
                })
                continue
            if found_parts:
                results.append({
                    'brand': brand,
                    'brand_number': brand_number,
                    'part_number': part_no,
                    'status': 'Found',
                    'car_count': len(found_parts),
                    'top_cars': [
                        f"{(p.brand or '—')} | {(p.part_number or '—')}"
                        for p in found_parts[:3]
                    ],
                    'in_basket': part_no in basket_part_numbers or any(
                        sanitize_part_number(p.part_number) in basket_part_numbers
                        for p in found_parts
                    ),
                })
            else:
                results.append({
                    'brand': brand,
                    'brand_number': brand_number,
                    'part_number': part_no,
                    'status': 'Not Found',
                    'car_count': 0,
                    'top_cars': [],
                    'in_basket': False,
                })
        return results

    @staticmethod
    def _build_summary(rows_data, parts_by_part_num, matched_car_ids, missed_rows, added_to_basket):
        found_count = 0
        for row_data in rows_data:
            if row_data.get('bulk_miss'):
                continue
            if (
                row_data['brand']
                and row_data['brand_number']
                and row_data['part_number'] in parts_by_part_num
            ):
                found_count += 1
        return {
            'searched_count': len(rows_data),
            'found_count': found_count,
            'car_models_matched': len(matched_car_ids),
            'missed_count': len(missed_rows),
            'added_to_basket': added_to_basket,
            'missed_rows': missed_rows,
        }

    @staticmethod
    def build_sample_workbook():
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = 'Bulk search sample'
        sheet.append(['Cross Brand', 'Cross Code', 'Part Number'])
        sheet.append(['AISIN', 'AS-12345', '3231A047'])
        return workbook
