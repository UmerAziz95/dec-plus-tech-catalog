"""
Format-agnostic spreadsheet loading for the import pipeline.

Import parsing code (see import_services.py) only ever calls a handful of
openpyxl APIs: workbook.sheetnames, workbook[name], workbook.close(), and
sheet.iter_rows(min_row=, max_row=, values_only=True). Wrapping .xls (via
xlrd) and .csv (via the stdlib csv module) behind that same tiny surface
means the rest of the import code doesn't need to know which format it's
reading.
"""
import csv
from pathlib import Path

import openpyxl

SUPPORTED_EXTENSIONS = ('.xlsx', '.xlsm', '.xls', '.csv')


class _RowsSheet:
    def __init__(self, rows):
        self._rows = rows

    def iter_rows(self, min_row=1, max_row=None, values_only=True):
        start = max(min_row - 1, 0)
        end = max_row if max_row is not None else len(self._rows)
        for row in self._rows[start:end]:
            yield tuple(row)


class _RowsWorkbook:
    def __init__(self, sheets):
        self._sheets = {name: _RowsSheet(rows) for name, rows in sheets.items()}
        self.sheetnames = list(self._sheets.keys())

    def __getitem__(self, name):
        return self._sheets[name]

    def close(self):
        pass


def _load_xls_workbook(path):
    import xlrd

    book = xlrd.open_workbook(str(path))
    sheets = {}
    for sheet in book.sheets():
        sheets[sheet.name] = [tuple(sheet.row_values(r)) for r in range(sheet.nrows)]
    return _RowsWorkbook(sheets)


def _read_csv_rows(path):
    last_error = None
    for encoding in ('utf-8-sig', 'utf-8', 'latin-1'):
        try:
            with open(path, newline='', encoding=encoding) as fh:
                sample = fh.read(65536)
                fh.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=',;\t|')
                except csv.Error:
                    dialect = csv.excel
                return [tuple(row) for row in csv.reader(fh, dialect)]
        except UnicodeDecodeError as exc:
            last_error = exc
            continue
    if last_error:
        raise ValueError('Could not decode CSV file (unsupported text encoding).') from last_error
    raise ValueError('Could not decode CSV file (unsupported text encoding).')


def _load_csv_workbook(path):
    # Import code resolves sheets by alias and falls back to the first sheet,
    # so a single-sheet CSV works for every import type (Cross Car + Cross Code).
    return _RowsWorkbook({'CSV': _read_csv_rows(path)})


def load_workbook(path):
    """Open .xlsx/.xlsm/.xls/.csv and return an openpyxl-Workbook-like object."""
    suffix = Path(path).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        allowed = ', '.join(SUPPORTED_EXTENSIONS)
        raise ValueError(f'Unsupported file type. Allowed: {allowed}')
    if suffix == '.csv':
        return _load_csv_workbook(path)
    if suffix == '.xls':
        return _load_xls_workbook(path)
    return openpyxl.load_workbook(path, data_only=True, read_only=True)
